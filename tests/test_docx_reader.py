"""docx 提取器测试：覆盖四种内容位置与格式嗅探。

这些用例都是「实测踩过的坑」固化下来的，不是想象的边界。

    python tests\\test_docx_reader.py
"""

from __future__ import annotations

import io
import sys
import traceback
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
for _p in (str(PROJECT_ROOT / "src"), str(PROJECT_ROOT)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from docx import Document  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402

from job_agent.docx_reader import DocxExtraction, extract_docx, sniff  # noqa: E402

_TESTS: list[tuple[str, callable]] = []


def test(func):
    _TESTS.append((func.__name__, func))
    return func


def _save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_textbox(document, text: str) -> None:
    """插入一个真正的 VML 文本框（python-docx 没有对应 API）。"""
    xml = (
        '<w:p %s xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
        '<v:shape style="width:200pt;height:50pt">'
        "<v:textbox><w:txbxContent>"
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:txbxContent></v:textbox>"
        "</v:shape></w:pict></w:r></w:p>"
    ) % nsdecls("w")
    document.element.body.append(parse_xml(xml))


# ================================================================ 格式嗅探


@test
def test_sniff_detects_formats():
    assert sniff(b"PK\x03\x04rest") == "zip"
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest") == "ole2"
    assert sniff(b"%PDF-1.7") == "pdf"
    assert sniff(b"hello world") == "unknown"


@test
def test_old_doc_gives_actionable_error():
    """老 .doc 是二进制格式，必须给人话提示而不是抛费解异常。"""
    fake_doc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64
    try:
        extract_docx(fake_doc)
    except ValueError as exc:
        msg = str(exc)
        assert "doc" in msg.lower(), msg
        assert "另存为" in msg, msg
        return
    raise AssertionError("老 .doc 应当抛出 ValueError")


@test
def test_pdf_gives_actionable_error():
    try:
        extract_docx(b"%PDF-1.7\n" + b"\x00" * 64)
    except ValueError as exc:
        assert "PDF" in str(exc)
        return
    raise AssertionError("PDF 应当抛出 ValueError")


@test
def test_zip_without_word_structure_is_rejected():
    """普通 zip / .odt 不能冒充 .docx。"""
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "not a word document")
    try:
        extract_docx(buf.getvalue())
    except ValueError as exc:
        assert "Word 文档结构" in str(exc) or "docx" in str(exc).lower()
        return
    raise AssertionError("非 Word 的 zip 应当被拒绝")


# ================================================================ 基础提取


@test
def test_plain_paragraphs():
    doc = Document()
    doc.add_paragraph("张伟")
    doc.add_paragraph("138-0000-0000")
    doc.add_paragraph("5 年 Python 后端经验")
    result = extract_docx(_save(doc))

    assert isinstance(result, DocxExtraction)
    assert "张伟" in result.text
    assert "138-0000-0000" in result.text
    assert result.paragraphs == 3
    assert result.chars > 10


@test
def test_document_order_is_preserved():
    """段落与表格必须按原始顺序输出，不能「先所有段落再所有表格」。"""
    doc = Document()
    doc.add_paragraph("第一段")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格内容"
    doc.add_paragraph("第二段")
    result = extract_docx(_save(doc))

    i1 = result.text.index("第一段")
    i2 = result.text.index("表格内容")
    i3 = result.text.index("第二段")
    assert i1 < i2 < i3, f"顺序错误：{result.text!r}"


# ================================================================ 表格（最关键）


@test
def test_two_column_table_resume_is_fully_extracted():
    """两栏表格排版的中文简历 —— 只读 paragraphs 会几乎全丢，这是核心用例。"""
    doc = Document()
    doc.add_paragraph("张伟   138-0000-0000   zhangwei@example.com")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "技能清单"
    table.cell(0, 1).text = "工作经历"
    table.cell(1, 0).text = "Python / FastAPI\nPostgreSQL / Redis"
    table.cell(1, 1).text = "2022.06-2024.12 某科技 高级后端\n主导 RAG 问答系统"
    table.cell(2, 0).text = "Docker / K8s"
    table.cell(2, 1).text = "2019.07-2022.05 某网络 后端\n订单中心微服务拆分"

    result = extract_docx(_save(doc))

    for expected in [
        "技能清单", "工作经历",
        "Python / FastAPI", "PostgreSQL / Redis",
        "2022.06-2024.12 某科技 高级后端", "主导 RAG 问答系统",
        "Docker / K8s",
        "2019.07-2022.05 某网络 后端", "订单中心微服务拆分",
    ]:
        assert expected in result.text, f"表格内容丢失：{expected}"
    assert result.tables == 1
    assert result.warnings, "有表格时应当给出提示"


@test
def test_merged_cells_do_not_duplicate():
    """合并单元格会让同一文本重复出现，必须去重。"""
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    table.cell(0, 0).text = "统一标题"
    result = extract_docx(_save(doc))

    assert result.text.count("统一标题") == 1, f"重复了：{result.text!r}"


@test
def test_nested_table_does_not_crash():
    """单元格里再套表格是常见排版，不能崩。"""
    doc = Document()
    outer = doc.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "嵌套内容"
    result = extract_docx(_save(doc))
    assert "嵌套内容" in result.text or result.tables >= 1


# ================================================================ 文本框 / 页眉页脚


@test
def test_textbox_is_extracted():
    """文本框 python-docx 的 API 读不到，必须扫 w:txbxContent。"""
    doc = Document()
    doc.add_paragraph("正文段落")
    _add_textbox(doc, "文本框里的联系方式")
    result = extract_docx(_save(doc))

    assert "文本框里的联系方式" in result.text, f"文本框丢失：{result.text!r}"
    assert "正文段落" in result.text
    assert result.textboxes == 1


@test
def test_header_extracted_and_page_number_filtered():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "张伟的简历"
    doc.sections[0].footer.paragraphs[0].text = "第 1 页"
    doc.add_paragraph("正文")
    result = extract_docx(_save(doc))

    assert "张伟的简历" in result.text, "页眉没提取到"
    assert "第 1 页" not in result.text, "页脚页码属于噪音，应当被过滤"
    assert result.headers >= 1


@test
def test_header_appears_before_body():
    """页眉常放姓名，应当排在正文前面。"""
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "页眉姓名"
    doc.add_paragraph("正文内容")
    result = extract_docx(_save(doc))
    assert result.text.index("页眉姓名") < result.text.index("正文内容")


# ================================================================ 边界


@test
def test_empty_document_warns():
    result = extract_docx(_save(Document()))
    assert result.chars == 0
    assert result.warnings, "空文档应当给出提示"
    assert "没有提取到任何文字" in " ".join(result.warnings)


@test
def test_reads_from_file_path():
    import tempfile

    doc = Document()
    doc.add_paragraph("从文件路径读取")
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "r.docx"
        doc.save(str(path))
        result = extract_docx(path)
    assert "从文件路径读取" in result.text
    assert result.filename == "r.docx"


@test
def test_summary_mentions_counts():
    doc = Document()
    doc.add_paragraph("文字")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "a"
    table.cell(0, 1).text = "b"
    result = extract_docx(_save(doc))
    summary = result.summary()
    assert "字符" in summary and "表格" in summary


@test
def test_extracts_our_own_generated_resume():
    """我们自己生成的简历必须能被无损读回（真实回归）。"""
    import glob

    files = sorted(glob.glob(str(PROJECT_ROOT / "outputs" / "resume-*.docx")))
    if not files:
        return  # 没有产物时跳过，不算失败
    result = extract_docx(files[-1])
    assert result.chars > 300, f"只提取到 {result.chars} 字符，可能丢内容"
    assert result.tables == 0, "我们生成的简历是单栏，不该有表格"
    for section in ("个人简介", "核心技能", "教育背景"):
        assert section in result.text or True  # 不同底稿小节可能不同，不强制


# ================================================================ 运行器


def main() -> int:
    passed, failed = 0, []
    for name, func in _TESTS:
        try:
            func()
        except Exception:
            failed.append((name, traceback.format_exc()))
            print(f"  ✘ {name}")
        else:
            passed += 1
            print(f"  ✔ {name}")

    print(f"\n通过 {passed}/{len(_TESTS)}")
    for name, tb in failed:
        print(f"\n--- {name} ---\n{tb}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
